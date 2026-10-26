"""Build the independent handbook from reviewed repository documentation."""
from pathlib import Path
import re,json
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_BREAK,WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'artifacts';OUT.mkdir(exist_ok=True)
FONT='Heiti SC'
doc=Document();
# Remove theme border residue from every paragraph style and use explicit fonts.
for st in doc.styles:
    pp=st.element.find(qn('w:pPr'))
    if pp is not None:
        for b in list(pp.findall(qn('w:pBdr'))):pp.remove(b)
section=doc.sections[0];section.page_width=Cm(21);section.page_height=Cm(29.7);section.top_margin=Cm(1.8);section.bottom_margin=Cm(1.8);section.left_margin=Cm(2);section.right_margin=Cm(2)
for name in ['Normal','Body Text','List Bullet','List Number']:
    st=doc.styles[name];st.font.name=FONT;st.font.size=Pt(10.5);st.element.rPr.rFonts.set(qn('w:eastAsia'),FONT);st.paragraph_format.line_spacing=1.25;st.paragraph_format.space_after=Pt(5)
for name,size in [('Title',24),('Heading 1',18),('Heading 2',13),('Heading 3',11)]:
    st=doc.styles[name];st.font.name=FONT;st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('000000');st.element.rPr.rFonts.set(qn('w:eastAsia'),FONT);st.paragraph_format.space_before=Pt(13);st.paragraph_format.space_after=Pt(8)
style=doc.styles.add_style('Code Block',1);style.font.name='Liberation Mono';style.font.size=Pt(8);style.element.rPr.rFonts.set(qn('w:eastAsia'),FONT);style.paragraph_format.line_spacing=1.05;style.paragraph_format.space_after=Pt(1)
header=section.header.paragraphs[0];header.text='G1 EDU 具身智能数据管线';header.style='Caption';header.runs[0].font.color.rgb=RGBColor.from_string('667085')
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('2026 年 10 月 25 日    ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

def hyperlink(p,text,url):
    rel=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rel);run=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'235E88');props.append(color);run.append(props);t=OxmlElement('w:t');t.text=text;run.append(t);link.append(run);p._p.append(link)

def inline(p,text):
    pattern=r'(\[[^\]]+\]\([^)]+\)|`[^`]+`|\*\*[^*]+\*\*)'
    for token in re.split(pattern,text):
        if not token:continue
        m=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',token)
        if m:hyperlink(p,m[1],m[2])
        elif token.startswith('`'):
            run=p.add_run(token[1:-1]);run.font.name='Liberation Mono';run.font.size=Pt(8.7)
        elif token.startswith('**'):p.add_run(token[2:-2]).bold=True
        else:p.add_run(token)

def para(text,style=None):
    p=doc.add_paragraph(style=style);inline(p,text);return p

def table(rows):
    rows=[x for x in rows if not all(re.fullmatch(r':?-+:?',c.strip()) for c in x)]
    t=doc.add_table(rows=1,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.style='Light Shading Accent 1';t.autofit=False
    count=len(rows[0]);widths={3:[2.7,6.9,7.4],4:[3.5,4.2,4.6,4.7],5:[3.0,3.5,3.5,3.5,3.5]}.get(count,[17/count]*count)
    # First column of bibliography is short and the source titles need more space.
    if rows[0][0].strip()=='编号':widths=[1.4,7.5,8.1]
    if rows[0][0].strip()=='参数':widths=[7.8,1.5,7.7]
    if rows[0][0].strip()=='数据':widths=[6.0,4.0,7.0]
    if rows[0][0].strip()=='结论':widths=[7.4,2.6,7.0]
    for i,width in enumerate(widths):t.columns[i].width=Cm(width)
    for ridx,row in enumerate(rows):
        cells=t.rows[0].cells if ridx==0 else t.add_row().cells
        for i,value in enumerate(row):
            cells[i].width=Cm(widths[i]);cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cells[i].paragraphs[0];p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4);p.paragraph_format.line_spacing=1.15
            inline(p,value.strip())
            for run in p.runs:run.font.size=Pt(9)
        trpr=t.rows[ridx]._tr.get_or_add_trPr();nosplit=OxmlElement('w:cantSplit');trpr.append(nosplit)
        if ridx==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

doc.add_paragraph('G1 EDU 具身智能数据管线\n原理与上手手册',style='Title')
para('报告准备与真机到货后的操作依据')
para('核查日期 2026 年 10 月 25 日    报告日期 2026 年 10 月 26 日')
para('目标任务是将箱子从桌 A 搬到桌 B。本文解释遥操作设备分工、数据流、原始记录、质量审查、转换与训练，并提供到货后的分步验收方法。当前没有真机，本次已经完成的实验是公开 G1 数据的小规模训练和固定机身仿真检查。')
para('当前设备仅确认 43 个自由度。29 本体加双手各 7 与 Dex3 配置相符，但这仍是待订单核实的推断。')
para('先读原理与讲稿，再按上手步骤执行。文中的官方命令是条件化操作模板；本地已执行实验与尚待真机验证的内容分别标明。')
doc.add_picture(str(ROOT/'artifacts/figures/02_pipeline_and_action_contract.png'),width=Cm(17))
para('独立文档包含六部分原理、采集步骤、转换与训练、仿真实验、讲稿和一手来源索引。完整源码快照与可运行程序保存在同名项目目录。')

files=['00_阅读顺序与今日边界.md','01_技术原理与搬箱任务.md','02_真机上手手册.md','03_数据质检转换与训练.md','04_仿真演示与算力计划.md','05_报告讲稿与问答.md','06_来源索引与事实核查.md']
for filename in files:
    
    if not filename.startswith('03'):doc.add_page_break()
    lines=(ROOT/'docs'/filename).read_text().splitlines();i=0;code=False
    while i<len(lines):
        line=lines[i]
        if line.startswith('```'):code=not code;i+=1;continue
        if code:
            p=doc.add_paragraph(style='Code Block');p.add_run(line or ' ');p.paragraph_format.keep_with_next=(i+1<len(lines) and not lines[i+1].startswith('```'));i+=1;continue
        if not line.strip():i+=1;continue
        if line.startswith('#'):
            level=len(line)-len(line.lstrip('#'));title=line.lstrip('#').strip();title=re.sub(r'[^\w\s\u4e00-\u9fff]',' ',title);doc.add_heading(title,level=min(level,3));i+=1
            if filename.startswith('01') and level==1:
                doc.add_picture(str(ROOT/'artifacts/figures/01_devices_and_dataflow.png'),width=Cm(17));para('图 1 默认 Host 部署下的应用数据通路')
            if filename.startswith('04') and title=='两段视频各自证明什么':
                doc.add_picture(str(ROOT/'artifacts/figures/03_training_results.png'),width=Cm(17))
            continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                rows.append(lines[i].strip().strip('|').split('|'));i+=1
            table(rows);continue
        if line.startswith('- '):para(line[2:],'List Bullet');i+=1;continue
        if re.match(r'^\d+\. ',line):para(line);i+=1;continue
        para(line);i+=1
for element in doc.element.xpath('.//w:rFonts'):
    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        element.attrib.pop(qn('w:'+attr),None)
for st in doc.styles:
    for element in st.element.iter(qn('w:rFonts')):
        for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
            element.attrib.pop(qn('w:'+attr),None)
out=OUT/'G1_EDU具身智能数据管线原理与上手手册.docx';doc.save(out)
print(out)
